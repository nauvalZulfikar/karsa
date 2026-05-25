"""Dump outline (BAB + numbered sections) from DOCX by content scan, not style."""
import sys, io, os, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

src = sys.argv[1]
print(f"### {os.path.basename(src)}")

doc = Document(src)
all_paras = []
for p in doc.paragraphs:
    all_paras.append(p.text)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                all_paras.append(p.text)

# Match BAB headings, X.Y, X.Y.Z, and standalone Indonesian section names
hdr_patterns = [
    re.compile(r'^\s*BAB\s+[IVX]+\b', re.IGNORECASE),
    re.compile(r'^\s*\d+\.\d+(\.\d+)?\s+\S'),                        # 1.1, 1.1.1
    re.compile(r'^(KATA PENGANTAR|DAFTAR ISI|DAFTAR TABEL|DAFTAR GAMBAR|LAMPIRAN|DAFTAR PUSTAKA|RINGKASAN EKSEKUTIF)$', re.IGNORECASE),
]

for i, txt in enumerate(all_paras):
    t = (txt or '').strip()
    if not t:
        continue
    for pat in hdr_patterns:
        if pat.match(t):
            print(f"  [{i:04d}] {t[:160]}")
            break

# image + size info
imgs = sum(1 for rel in doc.part.rels.values() if 'image' in rel.target_ref)
print(f"### IMAGES: {imgs}  PARAGRAPHS: {len(all_paras)}  TABLES: {len(doc.tables)}")
